import io
import pytest
import pymupdf
import docx
from fastapi.testclient import TestClient
from backend.app.main import app
from backend.app.services.document_parser import (
    parse_document,
    sanitize_filename,
    detect_file_type,
    validate_document_size,
    UnsupportedFormatError,
    EmptyDocumentError,
    ScannedDocumentError,
    CorruptDocumentError,
    DocumentSizeLimitError,
)


@pytest.fixture
def client():
    """FastAPI TestClient fixture."""
    with TestClient(app) as test_client:
        yield test_client


def create_in_memory_pdf(pages_text: list[str]) -> bytes:
    """Helper to generate a synthetic in-memory PDF document."""
    doc = pymupdf.open()
    for text in pages_text:
        page = doc.new_page()
        if text.strip():
            page.insert_text((50, 72), text)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_in_memory_docx(paragraphs: list[str] = None, table_rows: list[list[str]] = None) -> bytes:
    """Helper to generate a synthetic in-memory DOCX document."""
    doc = docx.Document()
    if paragraphs:
        for p in paragraphs:
            doc.add_paragraph(p)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r_idx, row in enumerate(table_rows):
            for c_idx, cell in enumerate(row):
                table.cell(r_idx, c_idx).text = cell
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


# =====================================================================
# Unit Tests: Document Parser Service
# =====================================================================

def test_sanitize_filename():
    assert sanitize_filename("../../../etc/passwd/resume.pdf") == "resume.pdf"
    assert sanitize_filename("my resume (1) [final].pdf") == "my resume _1_ _final_.pdf"
    assert sanitize_filename("") == "unnamed_document"


def test_validate_document_size_empty():
    with pytest.raises(EmptyDocumentError, match="empty"):
        validate_document_size(b"")


def test_validate_document_size_oversized():
    # Test with custom 1 MB limit
    oversized_data = b"x" * (2 * 1024 * 1024)
    with pytest.raises(DocumentSizeLimitError, match="exceeds maximum permitted limit"):
        validate_document_size(oversized_data, max_mb=1)


def test_detect_file_type_unsupported():
    with pytest.raises(UnsupportedFormatError, match="Unsupported file format"):
        detect_file_type("resume.png", b"fake-png-bytes")


def test_detect_file_type_header_mismatch():
    with pytest.raises(CorruptDocumentError, match="lacks a valid PDF header"):
        detect_file_type("resume.pdf", b"not-a-pdf-header")

    with pytest.raises(CorruptDocumentError, match="lacks a valid DOCX/ZIP header"):
        detect_file_type("resume.docx", b"not-a-docx-header")


def test_parse_valid_pdf_single_page():
    pdf_bytes = create_in_memory_pdf(["John Doe\nSoftware Engineer\nPython, FastAPI, SQL"])
    parsed = parse_document(pdf_bytes, "john_doe_resume.pdf")

    assert parsed.file_type == "pdf"
    assert parsed.page_count == 1
    assert "John Doe" in parsed.text
    assert "Python, FastAPI, SQL" in parsed.text
    assert parsed.word_count > 0
    assert parsed.character_count > 0


def test_parse_valid_pdf_multi_page():
    pdf_bytes = create_in_memory_pdf([
        "Page 1: Experience summary and key projects.",
        "Page 2: Education, certifications, and technical proficiencies."
    ])
    parsed = parse_document(pdf_bytes, "multi_page_resume.pdf")

    assert parsed.file_type == "pdf"
    assert parsed.page_count == 2
    assert "Page 1" in parsed.text
    assert "Page 2" in parsed.text


def test_parse_scanned_or_blank_pdf():
    # PDF with blank page (simulating scanned/image or empty text layer)
    pdf_bytes = create_in_memory_pdf(["   "])
    with pytest.raises(ScannedDocumentError, match="No extractable text found in PDF"):
        parse_document(pdf_bytes, "scanned_resume.pdf")


def test_parse_corrupt_pdf():
    with pytest.raises(CorruptDocumentError):
        parse_document(b"%PDF-corrupted-data-without-valid-structure", "corrupt.pdf")


def test_parse_valid_docx():
    docx_bytes = create_in_memory_docx(paragraphs=[
        "Jane Smith",
        "Frontend Developer with 4 years experience in React and TypeScript.",
    ])
    parsed = parse_document(docx_bytes, "jane_smith.docx")

    assert parsed.file_type == "docx"
    assert parsed.page_count >= 1
    assert "Jane Smith" in parsed.text
    assert "Frontend Developer" in parsed.text


def test_parse_docx_with_tables():
    docx_bytes = create_in_memory_docx(
        paragraphs=["Skills and Proficiencies:"],
        table_rows=[
            ["Skill", "Proficiency Level"],
            ["Python", "Advanced"],
            ["Docker", "Intermediate"],
        ]
    )
    parsed = parse_document(docx_bytes, "table_resume.docx")

    assert parsed.file_type == "docx"
    assert "Skills and Proficiencies:" in parsed.text
    assert "Python" in parsed.text
    assert "Advanced" in parsed.text
    assert "Docker" in parsed.text


def test_parse_empty_docx():
    docx_bytes = create_in_memory_docx(paragraphs=[])
    with pytest.raises(EmptyDocumentError, match="contains no extractable"):
        parse_document(docx_bytes, "empty.docx")


def test_parse_corrupt_docx():
    with pytest.raises(CorruptDocumentError):
        parse_document(b"PK\x03\x04corrupt_zip_stream", "corrupt.docx")


# =====================================================================
# API Integration Tests: POST /api/v1/resumes/parse
# =====================================================================

def test_api_parse_valid_pdf(client):
    pdf_bytes = create_in_memory_pdf(["Alice Johnson\nBackend Developer with Python & AWS expertise."])
    response = client.post(
        "/api/v1/resumes/parse",
        files={"file": ("alice_resume.pdf", pdf_bytes, "application/pdf")},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["filename"] == "alice_resume.pdf"
    assert data["file_type"] == "pdf"
    assert data["page_count"] == 1
    assert "Alice Johnson" in data["text"]
    assert data["character_count"] > 0
    assert data["word_count"] > 0


def test_api_parse_valid_docx(client):
    docx_bytes = create_in_memory_docx(paragraphs=["Bob Wilson\nData Engineer with Spark & SQL."])
    response = client.post(
        "/api/v1/resumes/parse",
        files={"file": ("bob_resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["filename"] == "bob_resume.docx"
    assert data["file_type"] == "docx"
    assert "Bob Wilson" in data["text"]


def test_api_parse_unsupported_format(client):
    response = client.post(
        "/api/v1/resumes/parse",
        files={"file": ("resume.png", b"\x89PNG\r\n\x1a\nfakeimage", "image/png")},
    )
    assert response.status_code == 400
    assert "Unsupported file format" in response.json()["detail"]


def test_api_parse_empty_file(client):
    response = client.post(
        "/api/v1/resumes/parse",
        files={"file": ("empty.pdf", b"", "application/pdf")},
    )
    assert response.status_code == 400
    assert "empty" in response.json()["detail"].lower()


def test_api_parse_scanned_pdf(client):
    pdf_bytes = create_in_memory_pdf(["   "])
    response = client.post(
        "/api/v1/resumes/parse",
        files={"file": ("scanned.pdf", pdf_bytes, "application/pdf")},
    )
    assert response.status_code == 422
    assert "scanned" in response.json()["detail"].lower()


def test_api_parse_corrupt_file(client):
    response = client.post(
        "/api/v1/resumes/parse",
        files={"file": ("broken.pdf", b"%PDF-broken-garbage-data", "application/pdf")},
    )
    assert response.status_code == 400
    assert "Failed to open PDF" in response.json()["detail"]
