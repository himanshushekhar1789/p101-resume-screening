import io
import re
from pathlib import Path
from typing import NamedTuple, Optional
import pymupdf
import docx
from backend.app.config import get_settings


class DocumentParsingError(Exception):
    """Base exception for document parsing failures."""
    pass


class UnsupportedFormatError(DocumentParsingError):
    """Raised when an uploaded document has an unsupported extension or MIME type."""
    pass


class EmptyDocumentError(DocumentParsingError):
    """Raised when an uploaded document contains 0 bytes or no readable content."""
    pass


class ScannedDocumentError(DocumentParsingError):
    """Raised when a document (e.g. PDF) is purely scanned/image-based with no extractable text."""
    pass


class CorruptDocumentError(DocumentParsingError):
    """Raised when a document is damaged, incomplete, or malformed."""
    pass


class DocumentSizeLimitError(DocumentParsingError):
    """Raised when an uploaded document exceeds the maximum permitted size."""
    pass


class ParsedDocument(NamedTuple):
    filename: str
    file_type: str
    page_count: int
    character_count: int
    word_count: int
    text: str


def sanitize_filename(filename: str) -> str:
    """Sanitize the uploaded filename to prevent directory traversal or script injection."""
    name = Path(filename).name
    # Strip dangerous characters, keep alphanumeric, dots, dashes, underscores
    sanitized = re.sub(r"[^\w\.\- ]", "_", name)
    return sanitized or "unnamed_document"


def validate_document_size(file_bytes: bytes, max_mb: Optional[int] = None) -> None:
    """Ensure the file size is within acceptable limits and not empty."""
    if len(file_bytes) == 0:
        raise EmptyDocumentError("The uploaded file is empty (0 bytes).")

    settings = get_settings()
    limit_mb = max_mb if max_mb is not None else settings.max_upload_size_mb
    max_bytes = limit_mb * 1024 * 1024

    if len(file_bytes) > max_bytes:
        raise DocumentSizeLimitError(
            f"File size exceeds maximum permitted limit of {limit_mb} MB "
            f"({len(file_bytes) / (1024 * 1024):.2f} MB provided)."
        )


def detect_file_type(filename: str, file_bytes: bytes) -> str:
    """
    Detect and validate the file type using extension and magic bytes header.
    Returns 'pdf' or 'docx'.
    """
    ext = Path(filename).suffix.lower()
    if ext not in [".pdf", ".docx"]:
        raise UnsupportedFormatError(
            f"Unsupported file format '{ext}'. Only .pdf and .docx documents are supported."
        )

    # Magic byte verification
    if ext == ".pdf":
        if not file_bytes.startswith(b"%PDF"):
            raise CorruptDocumentError("The file has a .pdf extension but lacks a valid PDF header.")
        return "pdf"

    if ext == ".docx":
        # DOCX files are zip archives and start with the PK zip header
        if not file_bytes.startswith(b"PK\x03\x04"):
            raise CorruptDocumentError("The file has a .docx extension but lacks a valid DOCX/ZIP header.")
        return "docx"

    raise UnsupportedFormatError(f"Unsupported file type: {ext}")


def parse_pdf(file_bytes: bytes, sanitized_name: str) -> ParsedDocument:
    """
    Extract text from a PDF file using PyMuPDF (fitz).
    Preserves reading order and raises ScannedDocumentError if no text is found.
    """
    try:
        doc = pymupdf.open(stream=file_bytes, filetype="pdf")
    except Exception as e:
        raise CorruptDocumentError(f"Failed to open PDF document: {str(e)}")

    try:
        page_count = len(doc)
        if page_count == 0:
            raise EmptyDocumentError("The PDF document contains no pages.")

        page_texts = []
        for page_num in range(page_count):
            try:
                page = doc[page_num]
                # sort=True preserves reasonable left-to-right, top-to-bottom reading order
                text = page.get_text("text", sort=True)
                if text and text.strip():
                    page_texts.append(text.strip())
            except Exception as e:
                # Handle individual malformed page gracefully
                continue

        combined_text = "\n\n".join(page_texts).strip()

        if not combined_text:
            raise ScannedDocumentError(
                "No extractable text found in PDF. The document may be scanned or image-based. "
                "Text-based PDF resumes are required (OCR is not supported in this phase)."
            )

        char_count = len(combined_text)
        word_count = len(combined_text.split())

        return ParsedDocument(
            filename=sanitized_name,
            file_type="pdf",
            page_count=page_count,
            character_count=char_count,
            word_count=word_count,
            text=combined_text,
        )
    finally:
        doc.close()


def parse_docx(file_bytes: bytes, sanitized_name: str) -> ParsedDocument:
    """
    Extract text from a DOCX document using python-docx.
    Extracts paragraphs and table contents in document flow order.
    """
    try:
        stream = io.BytesIO(file_bytes)
        doc = docx.Document(stream)
    except Exception as e:
        raise CorruptDocumentError(f"Failed to open DOCX document: {str(e)}")

    text_parts = []

    try:
        # Traverse elements in document body order to preserve natural reading order
        for element in doc.element.body:
            if element.tag.endswith("p"):
                p = docx.text.paragraph.Paragraph(element, doc)
                if p.text.strip():
                    text_parts.append(p.text.strip())
            elif element.tag.endswith("tbl"):
                table = docx.table.Table(element, doc)
                table_lines = []
                for row in table.rows:
                    # Collect cell text while avoiding duplicated text from merged cells
                    row_cells = []
                    seen_cells = set()
                    for cell in row.cells:
                        txt = cell.text.strip()
                        if txt and txt not in seen_cells:
                            seen_cells.add(txt)
                            row_cells.append(txt)
                    if row_cells:
                        table_lines.append(" | ".join(row_cells))
                if table_lines:
                    text_parts.append("\n".join(table_lines))

        combined_text = "\n\n".join(text_parts).strip()

        if not combined_text:
            raise EmptyDocumentError("The DOCX document contains no extractable paragraph or table text.")

        char_count = len(combined_text)
        word_count = len(combined_text.split())
        page_estimate = max(1, len(doc.sections))

        return ParsedDocument(
            filename=sanitized_name,
            file_type="docx",
            page_count=page_estimate,
            character_count=char_count,
            word_count=word_count,
            text=combined_text,
        )
    except DocumentParsingError:
        raise
    except Exception as e:
        raise CorruptDocumentError(f"Error while parsing DOCX content: {str(e)}")


def parse_document(file_bytes: bytes, filename: str) -> ParsedDocument:
    """
    High-level entry point for document parsing.
    Validates, detects file format, and delegates to the appropriate parser.
    """
    sanitized_name = sanitize_filename(filename)
    validate_document_size(file_bytes)
    file_type = detect_file_type(sanitized_name, file_bytes)

    if file_type == "pdf":
        return parse_pdf(file_bytes, sanitized_name)
    elif file_type == "docx":
        return parse_docx(file_bytes, sanitized_name)
    else:
        raise UnsupportedFormatError(f"Unsupported file type: {file_type}")
